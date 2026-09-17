"""Enriquecimento (rodar local/manual) — historico de vendas do ano a partir do
export XLSX da Hubla.

A Hubla nao tem API de listagem: o webhook n8n so grava a partir de 06/08/2026, e a
planilha viva ([MXP-FP01][2026][BACKUP]) comeca em 01/08. Todo o ano anterior a isso
(Le01 em fev, Le02 em mar-abr) so existe no export do painel.

Uso:  /usr/bin/python3 import_vendas_hubla.py ~/Downloads/<export>.xlsx
Gera `vendas_ano.json`, que o build.py usa como base do ano. A planilha viva continua
mandando nas vendas novas (merge por id_fatura, planilha ganha — e dela que vem o
status de reembolso, que o export nao traz).
"""
import json
import os
import sys
from datetime import datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "vendas_ano.json")

# o produto do evento. O export tambem traz "Palestras MXP 2025", que e outro
# produto (evento do ano passado) e nao entra na conta de 2026.
PRODUTO = "[MXP] Memorável Experience 2026"

COL = {
    "id": "ID da fatura",
    "status": "Status da fatura",
    "pago_em": "Data de pagamento",
    "criado_em": "Data de criação",
    "oferta": "Nome da oferta",
    "produto": "Nome do produto",
    "metodo": "Método de pagamento",
    "parcelas": "Parcelas",
    "valor_produto": "Valor do produto",
    "valor": "Valor total",          # mesma base do campo `valor` da planilha viva
    "liquido": "Valor Líquido",
    "src": "UTM Origem",
    "medium": "UTM Mídia",
    "camp": "UTM Campanha",
    "cont": "UTM Conteúdo",
}


def dstr(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    for f in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = os.path.expanduser(sys.argv[1])
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    idx = {h: i for i, h in enumerate(hdr)}
    falta = [c for c in COL.values() if c not in idx]
    if falta:
        sys.exit("colunas ausentes no export: %s" % falta)

    vendas, fora = [], 0
    for r in it:
        if not r[idx[COL["id"]]]:
            continue
        if (r[idx[COL["produto"]]] or "").strip() != PRODUTO:
            fora += 1
            continue
        d = dstr(r[idx[COL["pago_em"]]]) or dstr(r[idx[COL["criado_em"]]])
        vendas.append({
            "id": str(r[idx[COL["id"]]]).strip(),
            "d": d,
            "status": (r[idx[COL["status"]]] or "").strip().lower().replace("paga", "paid"),
            "oferta": (r[idx[COL["oferta"]]] or "").strip() or "(sem oferta)",
            "metodo": (r[idx[COL["metodo"]]] or "").strip(),
            "parcelas": int(float(r[idx[COL["parcelas"]]] or 1)),
            "valor": round(float(r[idx[COL["valor"]]] or 0), 2),
            "valor_produto": round(float(r[idx[COL["valor_produto"]]] or 0), 2),
            "liquido": round(float(r[idx[COL["liquido"]]] or 0), 2),
            "utm_source": (r[idx[COL["src"]]] or "").strip(),
            "utm_medium": (r[idx[COL["medium"]]] or "").strip(),
            "utm_campaign": (r[idx[COL["camp"]]] or "").strip(),
            "utm_content": (r[idx[COL["cont"]]] or "").strip(),
        })

    vendas.sort(key=lambda v: v["d"])
    periodo = [vendas[0]["d"], vendas[-1]["d"]] if vendas else ["", ""]
    out = {
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "origem": os.path.basename(src),
        "produto": PRODUTO,
        "periodo": periodo,
        "vendas": vendas,
    }
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))
    total = sum(v["valor"] for v in vendas)
    print("%d vendas | R$ %s | %s -> %s | %d linhas de outro produto ignoradas"
          % (len(vendas), format(total, ",.2f"), periodo[0], periodo[1], fora))
    print(OUT)


if __name__ == "__main__":
    main()
